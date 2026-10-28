import tempfile
import unittest
from pathlib import Path

from docx import Document

from word_to_markdown import convert_document, main


class ConvertDocumentTests(unittest.TestCase):
    def test_converts_headings_formatting_lists_and_tables(self):
        document = Document()
        document.add_heading("Overview", level=2)
        paragraph = document.add_paragraph()
        paragraph.add_run("Bold").bold = True
        paragraph.add_run(" and ")
        paragraph.add_run("italic").italic = True
        document.add_paragraph("First item", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "A|B"
        table.cell(1, 1).text = "One"

        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "source.docx"
            document.save(input_path)

            markdown = convert_document(input_path)

        self.assertEqual(
            markdown,
            "## Overview\n\n**Bold** and *italic*\n\n"
            "- First item\n\n| Name | Value |\n| --- | --- |\n| A\\|B | One |\n",
        )

    def test_cli_writes_markdown_to_requested_path(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "source.docx"
            output_path = Path(directory) / "converted.md"
            document = Document()
            document.add_paragraph("Hello")
            document.save(input_path)

            result = main([str(input_path), "--output", str(output_path)])

            self.assertEqual(result, 0)
            self.assertEqual(output_path.read_text(encoding="utf-8"), "Hello\n")

    def test_cli_rejects_output_path_that_would_overwrite_input(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "source.docx"
            Document().save(input_path)

            with self.assertRaises(SystemExit) as error:
                main([str(input_path), "--output", str(input_path)])

            self.assertEqual(error.exception.code, 2)
            self.assertTrue(input_path.exists())


if __name__ == "__main__":
    unittest.main()
